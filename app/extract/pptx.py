"""pptx 抽出: スライドテキスト・表・グラフデータ・図形塗り・ノート。

図形の塗り色（`{shape-fill:...}`）は「赤で強調されている項目」型の質問で効く。
グループ図形は再帰的に辿る。
"""
from __future__ import annotations

import zipfile
from pathlib import Path

from app.corpus.walk import CorpusFile
from app.extract.markup import (StyledRun, fill_tag, gfm_table,
                                normalize_color, render_runs)
from app.extract.models import ExtractedDoc, Page
from app.extract.office_media import extract_media, media_block
from app.extract.ooxml_chart import extract_charts

MSO_SHAPE_GROUP = 6
MSO_SHAPE_PICTURE = 13
MSO_FILL_SOLID = 1


def _run_style(run) -> StyledRun:
    f = run.font
    color = None
    try:
        if f.color is not None and f.color.type is not None:
            color = normalize_color(getattr(f.color, "rgb", None))
    except (AttributeError, TypeError, ValueError):
        color = None
    return StyledRun(text=run.text, bold=bool(f.bold), italic=bool(f.italic),
                     underline=bool(f.underline), color=color)


def _frame_text(tf) -> str:
    lines = []
    for para in tf.paragraphs:
        t = render_runs([_run_style(r) for r in para.runs])
        if t.strip():
            lines.append(("  " * para.level) + ("- " if para.level else "") + t)
    return "\n".join(lines)


def _chart_md(chart) -> str:
    try:
        cats = [str(c) for c in chart.plots[0].categories]
    except (AttributeError, IndexError, TypeError, ValueError):
        cats = []
    rows = [["系列"] + cats]
    try:
        for s in chart.series:
            rows.append([str(s.name) if s.name else "(無名)"]
                        + ["" if v is None else str(v) for v in s.values])
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        ctype = str(chart.chart_type)
    except (AttributeError, ValueError):
        ctype = "?"
    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text
    except (AttributeError, ValueError):
        pass
    head = f"[グラフ ({ctype}){': ' + title if title else ''}]"
    return head + "\n" + gfm_table(rows) if len(rows) > 1 \
        else head + " (データ抽出不可)"


def _shape_fill(shape) -> str | None:
    try:
        f = shape.fill
        if f.type != MSO_FILL_SOLID:
            return None
        return normalize_color(getattr(f.fore_color, "rgb", None))
    except (AttributeError, TypeError, ValueError, KeyError):
        return None


def _walk_shapes(shapes, parts: list[str], warns: list[str]) -> int:
    n_pics = 0
    for shape in shapes:
        try:
            stype = shape.shape_type
            if stype == MSO_SHAPE_GROUP:
                n_pics += _walk_shapes(shape.shapes, parts, warns)
                continue
            if getattr(shape, "has_chart", False):
                parts.append(_chart_md(shape.chart))
                continue
            if getattr(shape, "has_table", False):
                parts.append(gfm_table(
                    [[_frame_text(c.text_frame) for c in row.cells]
                     for row in shape.table.rows]))
                continue
            if stype == MSO_SHAPE_PICTURE:
                n_pics += 1
                continue
            if getattr(shape, "has_text_frame", False):
                t = _frame_text(shape.text_frame)
                if t.strip():
                    tag = fill_tag(_shape_fill(shape), kind="shape-fill")
                    parts.append(f"{t}\n{tag}" if tag else t)
        except (AttributeError, TypeError, ValueError, KeyError) as e:
            warns.append(f"shape処理スキップ: {type(e).__name__}: {e}")
            continue
    return n_pics


def extract_pptx(cf: CorpusFile, assets_dir: Path | None = None) -> ExtractedDoc:
    doc = ExtractedDoc(doc_id=cf.doc_id, relpath=cf.relpath, mount=cf.mount,
                       project=cf.project, category=cf.category,
                       filetype="pptx", md5=cf.md5)
    try:
        from pptx import Presentation
        prs = Presentation(cf.raw_path)
    except (OSError, KeyError, ValueError, zipfile.BadZipFile) as e:
        doc.error = f"pptx open failed: {type(e).__name__}: {e}"
        return doc

    warns: list[str] = []
    total_pics = 0
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = [f"## スライド {i}"]
        total_pics += _walk_shapes(slide.shapes, parts, warns)
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"（ノート: {notes}）")
        except (AttributeError, ValueError):
            pass
        body_chars = sum(len(p) for p in parts[1:])
        doc.pages.append(Page(no=i, text="\n\n".join(parts), label=f"slide{i}",
                              needs_vision=body_chars < 120))

    charts, cwarns = extract_charts(cf.raw_path)
    warns.extend(cwarns)
    media = extract_media(cf.raw_path, cf.doc_id, assets_dir)
    warns.extend(media.warnings)
    mb = "\n\n".join(charts + ([media_block(media)] if media.items else []))
    if mb:
        doc.pages.append(Page(no=len(doc.pages) + 1,
                              text="## 埋め込み画像・オブジェクト\n" + mb,
                              label="media",
                              needs_vision=any(m.readable_as_image
                                               for m in media.items),
                              assets=media.asset_ids))
    for w in dict.fromkeys(warns):
        doc.warn(w)
    doc.meta.update(n_slides=len(prs.slides), n_pictures=total_pics,
                    n_charts=len(charts),
                    n_media=len(media.items),
                    n_emf_parsed=sum(1 for m in media.items if m.emf_markdown))
    return doc
